"""Excel exports.

Three workbooks:

  providers.xlsx         - the deliverable: one row per active individual
                            provider, plus a "Departed" tab for anyone the
                            directory no longer lists
  excluded_records.xlsx  - everything kept out, and why
  summary.xlsx           - run statistics, duplicate review queue, and the
                            incremental-update audit trail

``providers.xlsx`` carries two columns beyond a bare provider record -
Location Count and Site Confidence - because "primary site of care" is only
trustworthy for single-site providers; see ``providermap.validator.site_confidence``.
"""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .database import Database

log = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")


def _style(ws: Worksheet, widths: list[int]) -> None:
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28


def export_providers(db: Database, out: Path) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Providers"

    headers = [
        "Provider Name",
        "Credentials",
        "Specialty",
        "Primary Site of Care",
        "Practice Name",
        "Address",
        "City",
        "State",
        "ZIP",
        "Phone",
        "Profile URL",
        "NPI",
        "Provider Type",
        "Location Count",
        "Site Confidence",
        "Hospital Affiliation",
        "Accepting New Patients",
        "Rating",
        "Rating Count",
        "Languages",
        "Insurance Accepted",
        "First Seen",
        "Last Seen",
    ]
    ws.append(headers)

    rows = db.conn.execute(
        """
        SELECT full_name, credentials, specialty, primary_site_of_care,
               practice_name, address, city, state, zip, phone, profile_url,
               npi, provider_type, location_count, site_confidence,
               hospital_affiliation, accepting_new_patients, rating,
               rating_count, languages, insurance_accepted,
               first_seen, last_seen
        FROM providers WHERE is_active = 1
        ORDER BY last_name, first_name, full_name
        """
    ).fetchall()
    for r in rows:
        values = [r[k] for k in r.keys()]  # noqa: SIM118 - sqlite3.Row, not a dict
        anp_idx = r.keys().index("accepting_new_patients")
        if values[anp_idx] is not None:
            values[anp_idx] = "Yes" if values[anp_idx] else "No"
        ws.append(values)

    conf_col = headers.index("Site Confidence") + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, conf_col).value in ("low", "medium"):
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).fill = _FLAG_FILL

    _style(
        ws,
        [28, 16, 26, 34, 34, 32, 18, 8, 10, 15, 58, 14, 24, 14, 14, 28, 16, 10, 12, 26, 30, 22, 22],
    )

    gone = db.conn.execute(
        """
        SELECT full_name, credentials, specialty, primary_site_of_care, npi,
               first_seen, last_seen
        FROM providers WHERE is_active = 0 ORDER BY last_seen DESC
        """
    ).fetchall()
    ws2 = wb.create_sheet("Departed")
    ws2.append(
        [
            "Provider Name",
            "Credentials",
            "Specialty",
            "Last Known Site of Care",
            "NPI",
            "First Seen",
            "Last Seen",
        ]
    )
    for r in gone:
        ws2.append([r[k] for k in r.keys()])  # noqa: SIM118 - sqlite3.Row, not a dict
    if not gone:
        ws2.append(
            ["No providers have left the directory since tracking began.", "", "", "", "", "", ""]
        )
    _style(ws2, [28, 16, 26, 36, 14, 22, 22])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    log.info("Wrote %s (%s active providers, %s departed)", out, len(rows), len(gone))
    return len(rows)


def export_excluded(db: Database, out: Path) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Excluded"
    ws.append(["Name", "URL", "Classification", "Reason Excluded", "NPI", "Date Found"])

    rows = db.conn.execute(
        "SELECT name, url, classification, reason_excluded, npi, date_found "
        "FROM excluded_records ORDER BY classification, name"
    ).fetchall()
    for r in rows:
        ws.append([r[k] for k in r.keys()])  # noqa: SIM118 - sqlite3.Row, not a dict
    _style(ws, [40, 58, 26, 52, 14, 22])

    ws2 = wb.create_sheet("By Classification")
    ws2.append(["Classification", "Count"])
    for r in db.conn.execute(
        "SELECT classification, COUNT(*) c FROM excluded_records "
        "GROUP BY classification ORDER BY c DESC"
    ):
        ws2.append([r["classification"], r["c"]])
    _style(ws2, [30, 12])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    log.info("Wrote %s (%s excluded)", out, len(rows))
    return len(rows)


def export_summary(db: Database, out: Path) -> None:
    c = db.counts()
    dupes = db.find_probable_duplicates()
    last = db.last_run()

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    metrics = [
        ("Source used", db.get_stat("source_used", "unknown")),
        ("Total records found (URLs discovered)", c["urls_total"]),
        ("Records processed", c["urls_done"]),
        ("Valid providers (active)", c["providers"]),
        ("Providers no longer in directory (inactive)", c["providers_inactive"]),
        ("Facilities / organizations excluded", c["excluded"]),
        ("Unique locations", c["locations"]),
        ("Providers with >1 location (primary is inferred)", c["multi_location"]),
        ("Probable duplicate name+city groups (review needed)", len(dupes)),
        ("Field changes recorded (all time)", c["changes_logged"]),
        ("Errors encountered", c["urls_error"]),
        ("Still pending", c["urls_pending"]),
        ("Last run id", last.run_id if last else "-"),
        ("Date completed", datetime.now(timezone.utc).isoformat(timespec="seconds")),
    ]
    for k, v in metrics:
        ws.append([k, v])
    _style(ws, [52, 34])

    ws2 = wb.create_sheet("Provider Types")
    ws2.append(["Provider Type", "Count"])
    for r in db.conn.execute(
        "SELECT provider_type, COUNT(*) c FROM providers GROUP BY provider_type ORDER BY c DESC"
    ):
        ws2.append([r["provider_type"], r["c"]])
    _style(ws2, [32, 12])

    ws3 = wb.create_sheet("Site Confidence")
    ws3.append(["Site Confidence", "Count", "Meaning"])
    meaning = {
        "high": "Single location, NPPES-confirmed identity. Safe to join.",
        "medium": "2-3 locations, or identity unconfirmed. Review before joining.",
        "low": "4+ locations. 'Primary site' is not meaningful for this provider.",
    }
    for r in db.conn.execute(
        "SELECT site_confidence, COUNT(*) c FROM providers GROUP BY site_confidence ORDER BY c DESC"
    ):
        ws3.append([r["site_confidence"], r["c"], meaning.get(r["site_confidence"], "")])
    _style(ws3, [18, 10, 66])

    ws4 = wb.create_sheet("Duplicate Review")
    ws4.append(["Full Name", "City", "State", "Count", "NPIs", "URLs"])
    for d in dupes:
        ws4.append([d["full_name"], d["city"], d["state"], d["n"], d["npis"], d["urls"]])
    if not dupes:
        ws4.append(["No same-name/same-city collisions found.", "", "", "", "", ""])
    _style(ws4, [30, 18, 8, 8, 30, 70])

    ws_ch = wb.create_sheet("Recent Changes")
    ws_ch.append(["Run", "Provider", "NPI", "Field", "Old Value", "New Value", "Changed At"])
    prev = db.conn.execute(
        "SELECT run_id FROM runs WHERE finished_at IS NOT NULL "
        "ORDER BY run_id DESC LIMIT 1 OFFSET 1"
    ).fetchone()
    since = prev["run_id"] if prev else 0
    changes = db.changes_since(since)
    for ch in changes[:5000]:
        ws_ch.append(
            [ch.run_id, ch.full_name, ch.npi, ch.field, ch.old_value, ch.new_value, ch.changed_at]
        )
    if not changes:
        ws_ch.append(["No changes recorded since the previous run.", "", "", "", "", "", ""])
    _style(ws_ch, [6, 26, 14, 22, 40, 40, 22])

    ws_rh = wb.create_sheet("Run History")
    ws_rh.append(
        [
            "Run",
            "Mode",
            "Source",
            "Started",
            "Finished",
            "Discovered",
            "New",
            "Changed",
            "Unchanged",
            "Excluded",
            "Errors",
        ]
    )
    for r in db.recent_runs(50):
        ws_rh.append(
            [
                r.run_id,
                r.mode,
                r.source_used,
                r.started_at,
                r.finished_at,
                r.urls_discovered,
                r.providers_new,
                r.providers_changed,
                r.providers_unchanged,
                r.excluded,
                r.errors,
            ]
        )
    _style(ws_rh, [6, 12, 14, 22, 22, 11, 8, 9, 10, 10, 8])

    ws5 = wb.create_sheet("Errors")
    ws5.append(["URL", "Attempts", "Error", "Last Tried"])
    for r in db.conn.execute(
        "SELECT url, attempts, error_message, timestamp FROM scrape_log "
        "WHERE status='error' ORDER BY attempts DESC LIMIT 2000"
    ):
        ws5.append([r["url"], r["attempts"], r["error_message"], r["timestamp"]])
    _style(ws5, [58, 10, 60, 22])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    log.info("Wrote %s", out)


def export_all(db: Database, outdir: str | Path) -> None:
    outdir = Path(outdir)
    export_providers(db, outdir / "providers.xlsx")
    export_excluded(db, outdir / "excluded_records.xlsx")
    export_summary(db, outdir / "summary.xlsx")
