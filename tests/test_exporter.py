"""Tests for `export_organizations` - the export counterpart to
`ingest-organizations` / `enrich-organizations` / `link-organizations`, none
of which write anything human-readable on their own. The other three
exporter functions (providers/excluded/summary) have no dedicated unit tests
in this project - only this new one is covered here.
"""

from __future__ import annotations

from openpyxl import load_workbook

from providermap.database import Database
from providermap.exporter import export_organizations
from providermap.models import Confidence, Organization, Provider


class TestExportOrganizations:
    def test_writes_one_row_per_organization(self, db: Database, tmp_path) -> None:  # type: ignore[no-untyped-def]
        oid, _ = db.upsert_organization(
            Organization(
                name="Banner Desert Medical Center",
                organization_type="hospital",
                city="Mesa",
                state="AZ",
                source="CMS",
                source_id="1",
            )
        )
        assert oid is not None
        db.set_organization_website(oid, "https://bannerhealth.com", Confidence.MEDIUM)

        out = tmp_path / "organizations.xlsx"
        count = export_organizations(db, out)
        assert count == 1
        assert out.exists()

        wb = load_workbook(out)
        assert "Organizations" in wb.sheetnames
        ws = wb["Organizations"]
        headers = [c.value for c in ws[1]]
        assert headers == [
            "Organization Name",
            "Type",
            "Address",
            "City",
            "State",
            "ZIP",
            "Phone",
            "Website",
            "Website Confidence",
            "Linked Providers",
            "Source",
            "Source ID",
            "Created Date",
            "Updated Date",
        ]
        row = [c.value for c in ws[2]]
        assert row[headers.index("Organization Name")] == "Banner Desert Medical Center"
        assert row[headers.index("Website")] == "https://bannerhealth.com"
        assert row[headers.index("Website Confidence")] == "medium"
        assert row[headers.index("Linked Providers")] == 0

    def test_linked_providers_count_reflects_provider_organizations(
        self, db: Database, tmp_path
    ) -> None:  # type: ignore[no-untyped-def]
        oid, _ = db.upsert_organization(
            Organization(name="Some Hospital", source="CMS", source_id="1")
        )
        pid, _ = db.upsert_provider(Provider(full_name="Dr. A", npi="1194013169"))
        assert oid is not None and pid is not None
        db.link_provider_organization(pid, oid, source="name_match")

        out = tmp_path / "organizations.xlsx"
        export_organizations(db, out)
        wb = load_workbook(out)
        ws = wb["Organizations"]
        headers = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        assert row[headers.index("Linked Providers")] == 1

    def test_website_coverage_sheet_buckets_by_confidence(
        self, db: Database, tmp_path
    ) -> None:  # type: ignore[no-untyped-def]
        oid1, _ = db.upsert_organization(
            Organization(name="Has High Confidence", source="CMS", source_id="1")
        )
        oid2, _ = db.upsert_organization(
            Organization(name="Has No Website", source="CMS", source_id="2")
        )
        assert oid1 is not None and oid2 is not None
        db.set_organization_website(oid1, "https://example.com", Confidence.HIGH)

        out = tmp_path / "organizations.xlsx"
        export_organizations(db, out)
        wb = load_workbook(out)
        ws = wb["Website Coverage"]
        buckets = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
        assert buckets.get("high") == 1
        assert buckets.get("none") == 1

    def test_empty_database_writes_a_valid_workbook_with_no_rows(
        self, db: Database, tmp_path
    ) -> None:  # type: ignore[no-untyped-def]
        out = tmp_path / "organizations.xlsx"
        count = export_organizations(db, out)
        assert count == 0
        wb = load_workbook(out)
        ws = wb["Organizations"]
        assert ws.max_row == 1  # header only
